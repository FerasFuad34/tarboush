"""
خدمة التصدير — Markdown / JSON / PDF عربي / DOCX

ملاحظات تقنية مهمة:
- PDF: نستخدم ReportLab مع خط Noto Naskh Arabic مضمّن (OFL) + arabic_reshaper + python-bidi.
  التشكيل (reshaping) وتحويل الاتجاه (bidi) ضروريان لأن ReportLab لا يدعم العربية أصلاً.
  نُعطّل ligatures في reshaper لتجنّب محارف غير موجودة في الخط (مثل ligature «الله» U+FDF2).
- التغليف يتم بالكلمات لا بالحروف، مع قياس عرض السطر بالخط نفسه، ومحاذاة لليمين.
- DOCX: نضبط اتجاه الفقرات إلى RTL عبر خاصية w:bidi في OOXML.
"""
import json
import os
import re
from typing import List

import arabic_reshaper
from bidi.algorithm import get_display

# ============ أدوات عربية مشتركة ============
_RESHAPER = arabic_reshaper.ArabicReshaper(
    configuration={"support_ligatures": False, "delete_harakat": False}
)

MONTHS_AR = [
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر",
]


def format_date_ar(dt) -> str:
    """تاريخ عربي مقروء بلا رموز قد لا يدعمها الخط: 19 أكتوبر 2026"""
    return f"{dt.day} {MONTHS_AR[dt.month - 1]} {dt.year}"


SUMMARY_TITLES = {
    "glimpse": "لمحة سريعة (50 كلمة)",
    "bullets": "ملخص النقاط",
    "page": "ملخص الصفحة",
    "detailed": "ملخص تفصيلي",
    "top10": "أهم 10 أفكار",
}


def ar(text: str) -> str:
    """تشكيل عربي + ترتيب ثنائي الاتجاه مع أساس RTL"""
    text = (text or "").replace("\x00", " ")
    try:
        return get_display(_RESHAPER.reshape(text), base_dir="R")
    except Exception:
        return text


def _wrap(text: str, font_name: str, font_size: float, max_width: float) -> List[str]:
    """تغليف النص بالكلمات حسب عرض السطر الفعلي بالخط المستخدم"""
    from reportlab.pdfbase.pdfmetrics import stringWidth

    lines: List[str] = []
    for raw_line in (text or "").split("\n"):
        raw_line = raw_line.strip()
        if not raw_line:
            lines.append("")
            continue
        words = raw_line.split(" ")
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if stringWidth(ar(candidate), font_name, font_size) <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                # كلمة أطول من السطر: نقسمها
                while stringWidth(ar(word), font_name, font_size) > max_width and len(word) > 1:
                    cut = len(word)
                    while cut > 1 and stringWidth(ar(word[:cut]), font_name, font_size) > max_width:
                        cut -= 1
                    lines.append(word[:cut])
                    word = word[cut:]
                current = word
        if current:
            lines.append(current)
    return lines


# ============ Markdown ============
PROVIDER_LABELS_AR = {
    "youtube": "يوتيوب",
    "wikipedia": "ويكيبيديا",
    "wikiversity": "ويكي الجامعة",
    "wikibooks": "ويكي الكتب",
    "arxiv": "arXiv",
}


def _topic_blocks(topics) -> List[dict]:
    """
    توحيد شكل المواضيع للتصدير. تقبل قائمة صفوف ORM أو قواميس
    (حتى لا يعتمد التصدير على بنية قاعدة البيانات).
    """
    out: List[dict] = []
    for t in topics or []:
        if isinstance(t, dict):
            title = t.get("title") or ""
            excerpt = t.get("excerpt") or ""
            resources = t.get("resources") or []
        else:
            title = getattr(t, "title", "") or ""
            excerpt = getattr(t, "excerpt", "") or ""
            resources = list(getattr(t, "resources", []) or [])
        srcs = []
        for r in resources:
            if isinstance(r, dict):
                srcs.append({
                    "provider": r.get("provider", ""),
                    "title": r.get("title", ""),
                    "url": r.get("url", ""),
                })
            else:
                srcs.append({
                    "provider": getattr(r, "provider", ""),
                    "title": getattr(r, "title", ""),
                    "url": getattr(r, "url", ""),
                })
        out.append({"title": title, "excerpt": excerpt, "resources": srcs})
    return out


def export_markdown(lecture, transcript, summaries, glossary, quiz, topics=None) -> str:
    md = f"# {lecture.title}\n\n"
    md += f"**التاريخ:** {lecture.created_at.strftime('%Y-%m-%d %H:%M')}  \n"
    md += f"**الملف الأصلي:** {lecture.original_filename}  \n"
    if lecture.duration_seconds:
        md += f"**المدة:** {lecture.duration_seconds // 60} دقيقة\n"
    md += "\n---\n\n"

    if summaries:
        md += "## الملخصات\n\n"
        for s in summaries:
            md += f"### {SUMMARY_TITLES.get(s.level, s.level)}\n\n{s.content}\n\n"

    topic_blocks = _topic_blocks(topics)
    if topic_blocks:
        md += "## الشرح والمصادر\n\n"
        for i, t in enumerate(topic_blocks, 1):
            md += f"### {i}. {t['title']}\n\n"
            if t["excerpt"]:
                md += f"> {t['excerpt']}\n\n"
            if t["resources"]:
                for r in t["resources"]:
                    label = PROVIDER_LABELS_AR.get(r["provider"], r["provider"])
                    md += f"- **{label}:** [{r['title']}]({r['url']})\n"
            else:
                md += "_لم نجد مصدراً مطابقاً لهذا الموضوع._\n"
            md += "\n"

    if glossary:
        md += "## قاموس المصطلحات\n\n| المصطلح | English | التعريف من المحاضرة |\n|---|---|---|\n"
        for g in glossary:
            md += f"| {g.term_ar} | {g.term_en or '—'} | {g.definition} |\n"
        md += "\n"

    if quiz and quiz.questions:
        md += "## الاختبار\n\n"
        for i, q in enumerate(quiz.questions, 1):
            md += f"### س{i}: {q['question']}\n\n"
            for j, opt in enumerate(q["options"]):
                md += f"- {'✅' if j == q['answer'] else '⬜'} {opt}\n"
            md += f"\n> **الشرح:** {q.get('explanation', '')}\n\n"
            if q.get("source_quote"):
                md += f"> **الاقتباس من المحاضرة:** «{q['source_quote']}»\n\n"

    if transcript and transcript.full_text:
        md += "## النص الكامل\n\n"
        md += transcript.full_text + "\n"

    return md


# ============ JSON ============
def export_json(lecture, transcript, summaries, glossary, quiz, topics=None) -> str:
    data = {
        "lecture": {
            "id": lecture.id,
            "title": lecture.title,
            "filename": lecture.original_filename,
            "created_at": str(lecture.created_at),
            "duration": lecture.duration_seconds,
            "word_count": transcript.word_count if transcript else 0,
        },
        "transcript": {
            "text": transcript.full_text if transcript else "",
            "segments": transcript.segments if transcript else [],
        },
        "summaries": [{"level": s.level, "content": s.content} for s in (summaries or [])],
        "glossary": [{"ar": g.term_ar, "en": g.term_en, "definition": g.definition} for g in (glossary or [])],
        "quiz": quiz.questions if quiz else [],
        "topics": [
            {
                "title": t["title"],
                "excerpt": t["excerpt"],
                "resources": [
                    {
                        "provider": r["provider"],
                        "provider_label": PROVIDER_LABELS_AR.get(r["provider"], r["provider"]),
                        "title": r["title"],
                        "url": r["url"],
                    }
                    for r in t["resources"]
                ],
            }
            for t in _topic_blocks(topics)
        ],
    }
    return json.dumps(data, ensure_ascii=False, indent=2)


# ============ PDF ============
def export_pdf(lecture, transcript, summaries, glossary, quiz, output_path: str, topics=None):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.lib.colors import HexColor
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas

        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))  # backend/
        regular = os.path.join(base_dir, "fonts", "NotoNaskhArabic-Regular.ttf")
        bold = os.path.join(base_dir, "fonts", "NotoNaskhArabic-Bold.ttf")

        font_regular, font_bold = "Helvetica", "Helvetica-Bold"
        if os.path.exists(regular):
            pdfmetrics.registerFont(TTFont("TarboushArabic", regular))
            font_regular = "TarboushArabic"
        if os.path.exists(bold):
            pdfmetrics.registerFont(TTFont("TarboushArabicBold", bold))
            font_bold = "TarboushArabicBold"

        width, height = A4
        margin = 18 * mm
        max_width = width - 2 * margin
        primary = HexColor("#7f1d1d")

        c = canvas.Canvas(output_path, pagesize=A4)
        c.setTitle(lecture.title)

        state = {"y": height - margin - 5}

        def new_page():
            footer()
            c.showPage()
            state["y"] = height - margin

        def need(space):
            if state["y"] - space < margin + 10:
                new_page()

        def heading(text, size=15):
            need(size * 2.2)
            c.setFillColor(primary)
            c.setFont(font_bold, size)
            c.drawRightString(width - margin, state["y"], ar(text))
            state["y"] -= size * 1.9
            c.setStrokeColor(HexColor("#e5e7eb"))
            c.setLineWidth(0.7)
            c.line(margin, state["y"] + 4, width - margin, state["y"] + 4)
            state["y"] -= 6
            c.setFillColor(HexColor("#1f2937"))

        def paragraph(text, size=10.5, leading=6.2, color=None, indent=0):
            c.setFont(font_regular, size)
            c.setFillColor(color or HexColor("#1f2937"))
            for line in _wrap(text, font_regular, size, max_width - indent):
                need(size + leading)
                if line:
                    c.drawRightString(width - margin - indent, state["y"], ar(line))
                state["y"] -= size + leading

        def bullets(texts, size=10.5):
            """قائمة نقطية: ننظف أي علامات جاهزة في النص ثم نرسم النقطة بأنفسنا"""
            c.setFont(font_regular, size)
            for item in texts:
                item = re.sub(r"^[\-•–\*\s]+", "", item).strip()
                if not item:
                    continue
                wrapped = _wrap(item, font_regular, size, max_width - 12)
                for i, line in enumerate(wrapped):
                    need(size + 6)
                    if i == 0:
                        c.setFillColor(primary)
                        c.circle(width - margin + 2, state["y"] + size * 0.32, 2, stroke=0, fill=1)
                        c.setFillColor(HexColor("#1f2937"))
                    if line:
                        c.drawRightString(width - margin - 12, state["y"], ar(line))
                    state["y"] -= size + 6
                state["y"] -= 1.5

        def footer():
            """تذييل الصفحة الحالية (يُستدعى قبل كل showPage)"""
            c.saveState()
            c.setFont(font_regular, 8)
            c.setFillColor(HexColor("#9ca3af"))
            c.drawCentredString(width / 2, 10 * mm, ar("طربوش — منصة المحاضرات الذكية"))
            c.restoreState()

        # ---------- الغلاف/العنوان ----------
        c.setFillColor(primary)
        c.setFont(font_bold, 20)
        title_lines = _wrap(lecture.title, font_bold, 20, max_width)
        for line in title_lines[:3]:
            c.drawRightString(width - margin, state["y"], ar(line))
            state["y"] -= 26
        state["y"] -= 2

        c.setFont(font_regular, 9.5)
        c.setFillColor(HexColor("#6b7280"))
        # كل جزء يُرسم منفصلاً حتى لا يختلط الرقم بالاتجاه (bidi)
        x_cursor = width - margin
        meta_parts = [format_date_ar(lecture.created_at)]
        if lecture.duration_seconds:
            meta_parts.append(f"{lecture.duration_seconds // 60} دقيقة")
        if transcript:
            meta_parts.append(f"{transcript.word_count} كلمة")
        for part in meta_parts:
            label = ar(part)
            c.drawRightString(x_cursor, state["y"], label)
            from reportlab.pdfbase.pdfmetrics import stringWidth
            w = stringWidth(label, font_regular, 9.5)
            x_cursor -= w + 6
            c.setFillColor(HexColor("#d1d5db"))
            c.drawRightString(x_cursor, state["y"], ar("•"))
            x_cursor -= 12
            c.setFillColor(HexColor("#6b7280"))
        state["y"] -= 20
        c.setFillColor(HexColor("#1f2937"))

        # ---------- الملخصات ----------
        if summaries:
            heading("الملخصات")
            for s in summaries:
                heading(SUMMARY_TITLES.get(s.level, s.level), size=12)
                content = s.content or ""
                if s.level == "bullets":
                    items = [ln for ln in content.split("\n") if ln.strip()]
                    bullets(items)
                elif s.level == "top10":
                    # قائمة مرقمة: نرسم الرقم على اليمين يدوياً ثم النص (أدق من bidi مع الأرقام)
                    items = [re.sub(r"^[\-•–\*\s]+|^\d+[\).]\s*", "", ln.strip())
                             for ln in content.split("\n") if ln.strip()]
                    for idx, item in enumerate(items, 1):
                        wrapped = _wrap(item, font_regular, 10.5, max_width - 18)
                        for i, line in enumerate(wrapped):
                            need(16)
                            if i == 0:
                                c.setFont(font_bold, 10.5)
                                c.setFillColor(primary)
                                c.drawRightString(width - margin, state["y"], f"{idx}.")
                                c.setFillColor(HexColor("#1f2937"))
                            c.setFont(font_regular, 10.5)
                            c.drawRightString(width - margin - 18, state["y"], ar(line))
                            state["y"] -= 16
                        state["y"] -= 2
                else:
                    paragraph(content)
                state["y"] -= 6

        # ---------- الشرح والمصادر ----------
        topic_blocks = _topic_blocks(topics)
        if topic_blocks:
            heading("الشرح والمصادر")
            for i, t in enumerate(topic_blocks, 1):
                need(40)
                c.setFont(font_bold, 11.5)
                c.setFillColor(primary)
                c.drawRightString(width - margin, state["y"], ar(f"{i}. {t['title']}"))
                state["y"] -= 17
                c.setFillColor(HexColor("#1f2937"))
                if t["excerpt"]:
                    paragraph(t["excerpt"], size=9.5, leading=5)
                for r in t["resources"]:
                    label = PROVIDER_LABELS_AR.get(r["provider"], r["provider"])
                    c.setFont(font_regular, 9)
                    c.setFillColor(HexColor("#6b7280"))
                    need(14)
                    c.drawRightString(width - margin - 10, state["y"], ar(f"[{label}] {r['title']}"))
                    state["y"] -= 12
                    c.setFillColor(HexColor("#1d4ed8"))
                    # الرابط لاتيني: يُرسم بخط لاتيني (Noto Naskh لا يحتوي حروفاً لاتينية)
                    need(12)
                    url = r["url"]
                    from reportlab.pdfbase.pdfmetrics import stringWidth
                    while stringWidth(url, "Helvetica", 7.5) > max_width - 14 and len(url) > 20:
                        url = url[:-1]
                    c.setFont("Helvetica", 7.5)
                    c.drawRightString(width - margin - 10, state["y"], url)
                    c.setFont(font_regular, 9)
                    state["y"] -= 14
                    c.setFillColor(HexColor("#1f2937"))
                if not t["resources"]:
                    paragraph("لم نجد مصدراً مطابقاً لهذا الموضوع.", size=9, leading=5,
                              color=HexColor("#92400e"))
                state["y"] -= 4
            state["y"] -= 4

        # ---------- القاموس ----------
        if glossary:
            heading("قاموس المصطلحات")
            for g in glossary:
                need(30)
                term = g.term_ar + (f" ({g.term_en})" if g.term_en else "")
                c.setFont(font_bold, 11)
                c.setFillColor(primary)
                c.drawRightString(width - margin, state["y"], ar(term))
                state["y"] -= 15
                c.setFillColor(HexColor("#1f2937"))
                paragraph(g.definition, size=10)
                state["y"] -= 4

        # ---------- الاختبار ----------
        if quiz and quiz.questions:
            heading("الاختبار")
            for i, q in enumerate(quiz.questions, 1):
                need(46)
                # رقم السؤال يُرسم منفصلاً على اليمين (يتجنب اختلاط الأرقام مع RTL)
                c.setFillColor(primary)
                c.setFont(font_bold, 11)
                c.drawRightString(width - margin, state["y"], f"{i}.")
                c.setFillColor(HexColor("#111827"))
                for line in _wrap(q["question"], font_bold, 11, max_width - 20):
                    need(18)
                    c.setFont(font_bold, 11)
                    c.drawRightString(width - margin - 20, state["y"], ar(line))
                    state["y"] -= 16

                c.setFont(font_regular, 10.5)
                for j, opt in enumerate(q.get("options", [])):
                    need(18)
                    is_correct = j == q.get("answer")
                    color = HexColor("#15803d") if is_correct else HexColor("#374151")
                    # علامة مرسومة هندسياً (الخط العربي لا يحتوي بالضرورة على ✓)
                    if is_correct:
                        c.setFillColor(color)
                        c.circle(width - margin - 4, state["y"] + 3.4, 3.4, stroke=0, fill=1)
                    else:
                        c.setStrokeColor(color)
                        c.setLineWidth(0.9)
                        c.circle(width - margin - 4, state["y"] + 3.4, 3.4, stroke=1, fill=0)
                    c.setFillColor(color)
                    c.setFont(font_regular, 10.5)
                    for k, line in enumerate(_wrap(opt, font_regular, 10.5, max_width - 32)):
                        need(16)
                        c.setFont(font_regular, 10.5)
                        c.drawRightString(width - margin - 14, state["y"], ar(line))
                        if k < 3:
                            state["y"] -= 15

                if q.get("explanation"):
                    c.setFillColor(HexColor("#6b7280"))
                    paragraph(f"الشرح: {q['explanation']}", size=9, leading=5)
                c.setFillColor(HexColor("#1f2937"))
                state["y"] -= 8

        # ---------- النص الكامل ----------
        if transcript and transcript.full_text:
            heading("النص الكامل")
            paragraph(transcript.full_text, size=9.5, leading=5.4)

        footer()
        c.save()
        return output_path

    except Exception as e:
        print(f"[Export PDF] فشل: {e}")
        # بديل مضمون: ملف نصي بنفس المحتوى بدل الانهيار
        txt_path = output_path.replace(".pdf", ".txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(export_markdown(lecture, transcript, summaries, glossary, quiz))
        return txt_path


# ============ DOCX ============
def _set_rtl(paragraph):
    """ضبط اتجاه الفقرة إلى RTL داخل ملف Word (خاصية w:bidi في OOXML)"""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def export_docx(lecture, transcript, summaries, glossary, quiz, output_path: str, topics=None):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    title = doc.add_heading(lecture.title, level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(
        f"التاريخ: {lecture.created_at.strftime('%Y-%m-%d')} | الملف: {lecture.original_filename}"
    )

    if summaries:
        doc.add_heading("الملخصات", level=2)
        for s in summaries:
            doc.add_heading(SUMMARY_TITLES.get(s.level, s.level), level=3)
            doc.add_paragraph(s.content or "")

    topic_blocks = _topic_blocks(topics)
    if topic_blocks:
        doc.add_heading("الشرح والمصادر", level=2)
        for i, t in enumerate(topic_blocks, 1):
            doc.add_heading(f"{i}. {t['title']}", level=3)
            if t["excerpt"]:
                doc.add_paragraph(t["excerpt"])
            for r in t["resources"]:
                label = PROVIDER_LABELS_AR.get(r["provider"], r["provider"])
                doc.add_paragraph(f"[{label}] {r['title']} — {r['url']}")
            if not t["resources"]:
                doc.add_paragraph("لم نجد مصدراً مطابقاً لهذا الموضوع.")

    if glossary:
        doc.add_heading("قاموس المصطلحات", level=2)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "المصطلح"
        hdr[1].text = "English"
        hdr[2].text = "التعريف من المحاضرة"
        for g in glossary:
            row = table.add_row().cells
            row[0].text = g.term_ar
            row[1].text = g.term_en or "—"
            row[2].text = g.definition

    if quiz and quiz.questions:
        doc.add_heading("الاختبار", level=2)
        for i, q in enumerate(quiz.questions, 1):
            doc.add_paragraph(f"س{i}: {q['question']}")
            for j, opt in enumerate(q["options"]):
                mark = "✅" if j == q["answer"] else "⬜"
                doc.add_paragraph(f"{mark} {opt}")
            doc.add_paragraph(f"الشرح: {q.get('explanation', '')}")

    if transcript and transcript.full_text:
        doc.add_heading("النص الكامل", level=2)
        doc.add_paragraph(transcript.full_text)

    for para in doc.paragraphs:
        _set_rtl(para)

    doc.save(output_path)
    return output_path
